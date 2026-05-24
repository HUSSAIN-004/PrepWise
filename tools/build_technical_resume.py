from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\PrepWise-AI\Hussain_M_Kanchwala_Technical_Resume.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9E2F3"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_text(paragraph, text, bold=False, italic=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    add_text(p, text.upper(), bold=True, color="1F4D78", size=11)
    p_border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E2F3")
    p_border.append(bottom)
    p._p.get_or_add_pPr().append(p_border)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    add_text(p, text, size=9.5)


def add_role(doc, title, subtitle, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, title, bold=True, size=10.3)
    add_text(p, f" | {subtitle}", italic=True, size=9.7, color="555555")
    for bullet in bullets:
        add_bullet(doc, bullet)


def add_project(doc, title, tech, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, title, bold=True, size=10.2)
    add_text(p, f" - {tech}", italic=True, size=9.5, color="555555")
    for bullet in bullets:
        add_bullet(doc, bullet)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.5)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.05

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(0)
add_text(title, "HUSSAIN M. KANCHWALA", bold=True, size=18, color="0B2545")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(2)
add_text(
    subtitle,
    "T.E. Artificial Intelligence & Data Science Engineering | Full-Stack, Data Engineering & AI Enthusiast",
    size=9.5,
    color="555555",
)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(6)
add_text(contact, "Pune, Maharashtra | hussainkanchwala06@gmail.com | +91 9272906151", size=9.2)

add_section_heading(doc, "Professional Summary")
p = doc.add_paragraph()
add_text(
    p,
    "AI & Data Science engineering student with hands-on full-stack development experience, strong programming fundamentals, and growing interest in data engineering, machine learning, and AI-enabled systems. Experienced in MERN applications, authentication, dashboards, REST APIs, MongoDB integrations, responsive interfaces, and data-backed workflows.",
    size=9.6,
)

add_section_heading(doc, "Education")
edu = doc.add_paragraph()
add_text(edu, "Bachelor of Engineering - Artificial Intelligence & Data Science (T.E.)", bold=True, size=10.2)
add_text(edu, " | PVG College of Engineering and Technology, Pune", size=9.5)
add_bullet(doc, "Average SGPA till date: 9.2 / 10")

add_section_heading(doc, "Technical Skills")
skills_table = doc.add_table(rows=4, cols=2)
skills_table.autofit = False
set_table_borders(skills_table)
rows = [
    ("Programming Languages", "C++, Java, Python, JavaScript, Ruby, SQL"),
    ("Web Technologies", "HTML, CSS, JavaScript, React, Node.js, Express.js, Tailwind CSS, REST APIs"),
    ("Databases & Tools", "MongoDB Atlas, SQL, GitHub, VS Code, Postman, Excel, Power BI"),
    ("Core Subjects", "OOPS, OS, DBMS, Computer Graphics, Web Development, Data Science, Machine Learning Basics, Pattern Recognition, ANN, AI, NLP, Cyber Security"),
]
for row, (label, value) in zip(skills_table.rows, rows):
    row.cells[0].width = Inches(1.75)
    row.cells[1].width = Inches(5.55)
    for cell in row.cells:
        set_cell_margins(cell)
    set_cell_shading(row.cells[0], "F2F4F7")
    row.cells[0].paragraphs[0].add_run(label).bold = True
    row.cells[1].paragraphs[0].add_run(value)

add_section_heading(doc, "Experience")
add_role(
    doc,
    "Full Stack Developer Intern",
    "WhiskerBond Pvt. Ltd. | Current",
    [
        "Developing MERN-based features for pet and veterinary service workflows, including appointment booking, pet profiles, emergency appointments, birthday tracking, payments, and email notifications through AWS SES.",
        "Building responsive interfaces, scalable APIs, MongoDB integrations, and secure backend features while collaborating through GitHub.",
    ],
)
add_role(
    doc,
    "MERN Stack Intern",
    "Scalefull Technologies",
    [
        "Worked across frontend and backend modules using MongoDB, Express.js, React, and Node.js.",
        "Contributed to API development, database integration, testing, documentation, and feature debugging.",
    ],
)

add_section_heading(doc, "Selected Projects")
add_project(
    doc,
    "PrepWise AI - Placement Preparation Platform",
    "React, Node.js, Express.js, MongoDB, Gemini API, OpenTDB API",
    [
        "Built a full-stack placement preparation platform with protected authentication, reusable sidebar navigation, dashboard analytics, DSA practice tracking, resume analysis, mock interviews, aptitude tests, help center, and admin dashboard.",
        "Integrated Gemini for resume intelligence and mock interview feedback, Open Trivia DB for aptitude questions, and MongoDB-backed progress tracking, recommendations, and admin reporting.",
    ],
)
add_project(
    doc,
    "Tuition Class Management System",
    "MERN Stack, MongoDB, Express.js, React, Node.js",
    [
        "Developed a management platform for tuition class operations including student records, batch/course handling, enquiries, fee tracking, and admin workflows.",
        "Designed clean dashboards and backend APIs to reduce manual administration and improve visibility into class operations.",
    ],
)
add_project(
    doc,
    "DocBook",
    "MongoDB, Express.js, React.js, Node.js",
    [
        "Developed a full-stack appointment booking application with user, doctor, and admin dashboards, authentication, doctor search, and booking workflows.",
    ],
)
add_project(
    doc,
    "Additional Projects",
    "JavaScript, Python, Machine Learning",
    [
        "Built User Enquiry Form, Netflix Clone, Amazon Clone, To-Do List Keeper, Tic Tac Toe, Stone Paper Scissors, and House Price Prediction projects to strengthen frontend, backend, CRUD, logic-building, and predictive analysis skills.",
    ],
)

add_section_heading(doc, "Strengths")
add_bullet(doc, "Strong analytical and problem-solving skills with good ownership of assigned work.")
add_bullet(doc, "Quick learner with adaptability to new technologies and project requirements.")
add_bullet(doc, "Good understanding of full-stack development workflow from UI to APIs and database design.")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "Hussain M. Kanchwala | Technical Resume", size=8, color="777777")

doc.save(OUT)
print(OUT)
